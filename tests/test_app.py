"""End-to-end: Flask API -> job engine -> HTTP stack -> fake storefront -> SQLite -> Excel."""

import io
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.conftest import BASE, ISBNS, Book, FakeAwesomeBooks, build_env, make_books

SECTION = "https://www.awesomebooks.com/books/category/303/test-fiction"
# The v1.0 layouts the client's files already follow (v1.1 only appends columns after these).
V10_CATALOG_HEADERS = [
    "ISBN-13",
    "Title",
    "Author",
    "Format",
    "Availability",
    "New Available",
    "New Price (USD)",
    "New Qty",
    "Used Available",
    "Used Price (USD)",
    "Used Qty",
    "RRP (USD)",
    "Publisher",
    "Publication Date",
    "Pages",
    "Dimensions",
    "Weight",
    "Category",
    "Note",
    "Product URL",
    "Image URL",
]
V10_INVENTORY_HEADERS = [
    "SKU",
    "ISBN-13",
    "Status",
    "Quantity",
    "New Available",
    "New Price (USD)",
    "New Qty",
    "Used Available",
    "Used Price (USD)",
    "Used Qty",
    "RRP (USD)",
    "Title",
    "Author",
    "Format",
    "Listed on AwesomeBooks",
    "Note",
    "Product URL",
    "Checked At",
]


def wait_done(client, kind, job_id, timeout=30):
    deadline = time.time() + timeout
    p = None
    while time.time() < deadline:
        p = client.get(f"/api/{kind}/jobs/{job_id}").get_json()
        if not p["running"] and p["status"] in ("done", "failed", "paused"):
            return p
        time.sleep(0.05)
    raise AssertionError(f"job did not finish: {p}")


def rows(client, kind, job_id):
    return client.get(f"/api/{kind}/jobs/{job_id}/rows?limit=200").get_json()["rows"]


def sheet(client, kind, job_id):
    ws = load_workbook(io.BytesIO(client.get(f"/api/{kind}/jobs/{job_id}/export").data)).active
    headers = [c.value for c in ws[1]]
    return headers, [dict(zip(headers, [c.value for c in r], strict=False)) for r in ws.iter_rows(min_row=2)], ws


def start_catalog(c, url=SECTION, **extra):
    res = c.post("/api/catalog/jobs", json={"url": url, "sessions": 3, "pace": "fast", **extra})
    assert res.status_code == 201, res.get_json()
    return res.get_json()["job"]["id"]


# -- basics ---------------------------------------------------------------------------------


def test_health_config_and_dashboard(app_env):
    app, *_ = app_env
    c = app.test_client()
    assert c.get("/api/health").get_json()["ok"] is True
    cfg = c.get("/api/config").get_json()
    assert (cfg["ship_to"], cfg["currency"]) == ("US", "USD")
    assert {s["label"] for s in cfg["quick_sections"]} >= {"Used", "Bestsellers", "Bargain"}
    page = c.get("/")
    assert page.status_code == 200 and b"AwesomeBooks" in page.data
    assert "default-src 'self'" in page.headers["Content-Security-Policy"]


# -- catalog --------------------------------------------------------------------------------


def test_catalog_end_to_end(app_env, tmp_path):
    app, manager, site = app_env
    c = app.test_client()
    job_id = start_catalog(c)
    p = wait_done(c, "catalog", job_id)
    assert p["status"] == "done", p
    assert p["total"] == 4 and p["counts"]["done"] == 4  # in-stock books listed in the section
    assert p["stock"] == {"in_stock": 4, "out_of_stock": 0, "used": 3, "new": 3, "units": 1482 + 3 + 990 + 7 + 4 + 1}
    by_isbn = {r["isbn"]: r for r in rows(c, "catalog", job_id)}
    both = by_isbn[ISBNS[0]]
    assert (both["new_available"], both["new_price"], both["new_qty"]) == (1, 13.01, 1482)
    assert (both["used_available"], both["used_price"], both["used_qty"]) == (1, 8.84, 3)
    assert by_isbn[ISBNS[2]]["new_available"] == 0 and by_isbn[ISBNS[2]]["used_price"] == 10.57
    assert by_isbn[ISBNS[1]]["used_available"] == 0 and by_isbn[ISBNS[1]]["used_qty"] == 0
    # Every price page a visitor saw before switching to US was discarded, never stored.
    assert all(r["new_price"] in (None, 13.01, 18.74, 25.48) for r in by_isbn.values())

    headers, data, ws = sheet(c, "catalog", job_id)
    assert headers[:5] == ["ISBN-13", "Title", "Author", "Format", "Availability"]
    assert headers[5:11] == ["New Available", "New Price (USD)", "New Qty", "Used Available", "Used Price (USD)", "Used Qty"]
    # v1.1 adds the delivery windows at the end: every v1.0 column keeps its exact position.
    assert headers[:-2] == V10_CATALOG_HEADERS and headers[-2:] == ["New Delivery", "Used Delivery"]
    assert "ISBN10" not in " ".join(headers) and "ISBN-10" not in headers  # the client wants ISBN-13 only
    row = next(r for r in data if r["ISBN-13"] == ISBNS[0])
    assert (row["New Available"], row["New Price (USD)"], row["New Qty"]) == ("Yes", 13.01, 1482)
    assert (row["Used Available"], row["Used Price (USD)"], row["Used Qty"]) == ("Yes", 8.84, 3)
    assert ws["G1"].comment is not None and "US shopper" in ws["G1"].comment.text
    # The finished file was saved to "Downloads" by itself.
    job = manager.catalog.get_job(job_id)
    assert job["export_path"] and Path(job["export_path"]).is_file()


def test_catalog_quick_mode_uses_listing_data_only(app_env):
    app, *_ = app_env
    c = app.test_client()
    job_id = start_catalog(c, with_details=False)
    p = wait_done(c, "catalog", job_id)
    assert p["status"] == "done" and p["total"] == 4
    headers, data, _ = sheet(c, "catalog", job_id)
    assert "New Qty" not in headers and "Publisher" not in headers and "New Price (USD)" in headers
    assert "New Delivery" not in headers  # listing pages quote no delivery window
    row = next(r for r in data if r["ISBN-13"] == ISBNS[2])
    assert row["Used Price (USD)"] == 10.57 and row["New Available"] is None  # unknown, never "No"


def test_catalog_max_books_limit(app_env):
    app, *_ = app_env
    c = app.test_client()
    job_id = start_catalog(c, max_products="2")
    p = wait_done(c, "catalog", job_id)
    assert p["status"] == "done" and p["total"] == 2


def test_section_bigger_than_the_listing_window_is_sliced(tmp_path, monkeypatch):
    from awesomebooks_pro import scan

    shelf = make_books(150)
    site = FakeAwesomeBooks(shelf)
    site.page_size = 5
    site.page_cap = 4  # the "site" pages through 20 results per listing query
    monkeypatch.setattr(scan, "WINDOW", 20)
    monkeypatch.setattr(scan, "MAX_PAGES", 4)
    monkeypatch.setattr("awesomebooks_pro.parsing.LISTING_PAGE_SIZE", 5)
    monkeypatch.setattr("awesomebooks_pro.scan.LISTING_PAGE_SIZE", 5)
    app, manager, _ = build_env(tmp_path, site)
    c = app.test_client()
    job_id = start_catalog(c, with_details=False)
    p = wait_done(c, "catalog", job_id, timeout=60)
    assert p["status"] == "done", p
    in_stock = {b.isbn for b in shelf if b.offers()}
    found = {r["isbn"] for r in rows(c, "catalog", job_id)} | {r["isbn"] for r in _all_rows(c, job_id, p["total"])}
    assert found == in_stock  # every book reached although one listing shows only 20
    events = " ".join(e["message"] for e in manager.catalog.events(job_id, 5000))
    assert "parts by condition" in events and "parts by sub-category" in events
    assert len(in_stock) > 20 and "Discovery complete" in events


def test_limit_within_the_window_reads_the_first_books_in_site_order(tmp_path, monkeypatch):
    from awesomebooks_pro import scan

    shelf = make_books(150)
    site = FakeAwesomeBooks(shelf)
    site.page_size = 5
    site.page_cap = 4
    monkeypatch.setattr(scan, "WINDOW", 20)
    monkeypatch.setattr(scan, "MAX_PAGES", 4)
    monkeypatch.setattr("awesomebooks_pro.parsing.LISTING_PAGE_SIZE", 5)
    monkeypatch.setattr("awesomebooks_pro.scan.LISTING_PAGE_SIZE", 5)
    app, manager, _ = build_env(tmp_path, site)
    c = app.test_client()
    job_id = start_catalog(c, with_details=False, max_products=12)
    p = wait_done(c, "catalog", job_id)
    assert p["status"] == "done" and p["total"] == 12
    first_twelve = [b.isbn for b in shelf if b.offers()][:12]
    assert [r["isbn"] for r in rows(c, "catalog", job_id)] == first_twelve
    events = " ".join(e["message"] for e in manager.catalog.events(job_id, 500))
    assert "parts by" not in events  # no slicing requests were spent


def _all_rows(c, job_id, total):
    out = []
    for offset in range(0, total, 200):
        out += c.get(f"/api/catalog/jobs/{job_id}/rows?limit=200&offset={offset}").get_json()["rows"]
    return out


def test_bestsellers(app_env):
    app, *_ = app_env
    c = app.test_client()
    job_id = start_catalog(c, url="https://www.awesomebooks.com/books/bestsellers")
    p = wait_done(c, "catalog", job_id)
    assert p["status"] == "done" and p["total"] == 4
    assert {r["isbn"]: r["new_price"] for r in rows(c, "catalog", job_id)}[ISBNS[1]] == 18.74


def test_catalog_rejects_bad_links_and_missing_sections(app_env):
    app, *_ = app_env
    c = app.test_client()
    assert c.post("/api/catalog/jobs", json={"url": "https://example.com/x"}).status_code == 400
    assert c.post("/api/catalog/jobs", json={"url": f"{BASE}/book/{ISBNS[0]}/x"}).status_code == 400
    job_id = start_catalog(c, url="https://www.awesomebooks.com/books/category/999/nothing-here")
    p = wait_done(c, "catalog", job_id)
    assert p["status"] == "failed" and "does not exist" in p["job"]["error"]


def test_throttling_is_absorbed_not_recorded(app_env):
    app, manager, site = app_env
    site.fail_plan = [429, 503, 403, 429]
    c = app.test_client()
    job_id = start_catalog(c)
    p = wait_done(c, "catalog", job_id)
    assert p["status"] == "done" and p["counts"]["error"] == 0
    assert p["traffic"]["warnings"] >= 3
    assert {r["isbn"]: r["new_qty"] for r in rows(c, "catalog", job_id)}[ISBNS[0]] == 1482


def test_locale_lost_mid_run_never_leaks_uk_prices(app_env):
    app, manager, site = app_env
    site.forget_locale_on = set(range(3, 40, 4))  # sessions keep "forgetting" the country
    c = app.test_client()
    job_id = start_catalog(c)
    p = wait_done(c, "catalog", job_id)
    assert p["status"] == "done"
    prices = {r["isbn"]: (r["new_price"], r["used_price"]) for r in rows(c, "catalog", job_id)}
    assert prices[ISBNS[0]] == (13.01, 8.84) and prices[ISBNS[2]] == (None, 10.57) and prices[ISBNS[4]] == (25.48, 11.0)


def test_stop_then_resume_finishes_without_duplicates(tmp_path):
    from awesomebooks_pro.net import AdaptiveRateLimiter

    site = FakeAwesomeBooks(make_books(30))
    app, manager, _ = build_env(tmp_path, site)
    manager.limiter_factory = lambda interval: AdaptiveRateLimiter(0.03, jitter=0.0)
    c = app.test_client()
    job_id = start_catalog(c)
    time.sleep(0.4)
    assert c.post("/api/stop").get_json()["stopping"] is True
    p = wait_done(c, "catalog", job_id)
    assert p["status"] == "paused"
    assert c.post(f"/api/catalog/jobs/{job_id}/start", json={}).status_code == 200
    p = wait_done(c, "catalog", job_id, timeout=60)
    assert p["status"] == "done"
    isbns = [r["isbn"] for r in _all_rows(c, job_id, p["total"])]
    assert len(isbns) == len(set(isbns)) == sum(1 for b in site.books.values() if b.offers())


# -- stock monitor ----------------------------------------------------------------------------


def upload(c, text, name="list.csv"):
    res = c.post("/api/inventory/uploads", data={"file": (io.BytesIO(text.encode()), name)}, content_type="multipart/form-data")
    assert res.status_code == 201, res.get_json()
    return res.get_json()


def test_stock_check_end_to_end(app_env):
    app, manager, site = app_env
    c = app.test_client()
    isbn10 = ISBNS[1][3:12]
    total = sum((10 - i) * int(d) for i, d in enumerate(isbn10)) % 11
    isbn10 += "X" if (11 - total) % 11 == 10 else str((11 - total) % 11)
    csv = "\n".join(["ISBN", ISBNS[0], isbn10, f"9053-{ISBNS[2]}", ISBNS[3], ISBNS[5], "not-an-isbn", ISBNS[0]])
    up = upload(c, csv)
    assert up["upload"]["valid"] == 6 and up["upload"]["invalid"] == 1 and up["upload"]["duplicates"] == 1
    job_id = up["job"]["id"]
    assert c.post(f"/api/inventory/jobs/{job_id}/start", json={"sessions": 2, "pace": "fast"}).status_code == 200
    p = wait_done(c, "inventory", job_id)
    assert p["status"] == "done", p
    assert p["stock"]["in_stock"] == 4 and p["stock"]["out_of_stock"] == 2 and p["stock"]["not_found"] == 1

    headers, data, _ = sheet(c, "inventory", job_id)
    assert headers[:4] == ["SKU", "ISBN-13", "Status", "Quantity"]
    assert headers[:-2] == V10_INVENTORY_HEADERS and headers[-2:] == ["New Delivery", "Used Delivery"]
    by_sku = {r["SKU"]: r for r in data}
    both = by_sku[ISBNS[0]]
    assert (both["Status"], both["Quantity"], both["New Qty"], both["Used Qty"]) == ("In Stock", 1485, 1482, 3)
    assert (both["New Price (USD)"], both["Used Price (USD)"]) == (13.01, 8.84)
    assert by_sku[isbn10]["ISBN-13"] == ISBNS[1] and by_sku[isbn10]["Quantity"] == 990  # ISBN-10 checked as its ISBN-13
    assert by_sku[f"9053-{ISBNS[2]}"]["Used Available"] == "Yes" and by_sku[f"9053-{ISBNS[2]}"]["New Available"] == "No"
    assert (by_sku[ISBNS[3]]["Status"], by_sku[ISBNS[3]]["Quantity"], by_sku[ISBNS[3]]["Listed on AwesomeBooks"]) == (
        "Out of Stock",
        0,
        "Yes",
    )
    assert (by_sku[ISBNS[5]]["Status"], by_sku[ISBNS[5]]["Listed on AwesomeBooks"]) == ("Out of Stock", "No")
    assert by_sku["not-an-isbn"]["Status"] == "Invalid ISBN"
    assert [r["SKU"] for r in data] == [ISBNS[0], isbn10, f"9053-{ISBNS[2]}", ISBNS[3], ISBNS[5], "not-an-isbn", ISBNS[0]]
    assert site.search_calls == 5  # the duplicate ISBN was answered from the first check


def test_upload_rejects_unusable_files(app_env):
    app, *_ = app_env
    c = app.test_client()
    res = c.post(
        "/api/inventory/uploads", data={"file": (io.BytesIO(b"hello\nworld"), "x.csv")}, content_type="multipart/form-data"
    )
    assert res.status_code == 400 and "No valid ISBNs" in res.get_json()["error"]


def test_one_job_at_a_time(tmp_path):
    from awesomebooks_pro.net import AdaptiveRateLimiter

    site = FakeAwesomeBooks(make_books(20))
    app, manager, _ = build_env(tmp_path, site)
    manager.limiter_factory = lambda interval: AdaptiveRateLimiter(0.05, jitter=0.0)
    c = app.test_client()
    start_catalog(c)
    second = c.post("/api/catalog/jobs", json={"url": SECTION})
    assert second.status_code == 409
    c.post("/api/stop")
    manager.wait(10)


# -- security & export safety -------------------------------------------------------------------


def test_csrf_and_dns_rebinding_guards(app_env):
    app, *_ = app_env
    raw = app.test_client()
    raw.__class__ = raw.__class__.__mro__[1]  # a plain FlaskClient: no dashboard header
    assert raw.post("/api/stop").status_code == 403
    c = app.test_client()
    assert c.get("/api/config", headers={"Host": "evil.example"}).status_code == 403
    assert c.post("/api/stop", headers={"Origin": "http://evil.example"}).status_code == 403


def test_scraped_text_can_never_become_an_excel_formula(tmp_path):
    site = FakeAwesomeBooks([Book(ISBNS[0], "=HYPERLINK(evil)", new=(5.0, 2))])
    app, *_ = build_env(tmp_path, site)
    c = app.test_client()
    job_id = start_catalog(c)
    wait_done(c, "catalog", job_id)
    _, _, ws = sheet(c, "catalog", job_id)
    cell = ws["B2"]
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"


@pytest.mark.parametrize("value, expected", [("", None), ("0", None), ("1,500", 1500), ("abc", None), (25, 25)])
def test_limit_parsing(value, expected):
    from awesomebooks_pro.web import _limit

    assert _limit(value) == expected


def test_delivery_windows_reach_both_excel_files(app_env):
    app, manager, site = app_env
    site.books[ISBNS[4]].new_delivery = "Mon, November 9 - Thu, November 12"  # e.g. a pre-order
    c = app.test_client()
    job_id = start_catalog(c)
    assert wait_done(c, "catalog", job_id)["status"] == "done"
    _, data, ws = sheet(c, "catalog", job_id)
    by_isbn = {r["ISBN-13"]: r for r in data}
    assert (by_isbn[ISBNS[0]]["New Delivery"], by_isbn[ISBNS[0]]["Used Delivery"]) == (
        "Thu, November 5 - Tue, November 10",
        "Tue, November 3 - Fri, November 6",
    )
    assert by_isbn[ISBNS[1]]["Used Delivery"] is None  # New only: no Used window
    assert by_isbn[ISBNS[2]]["New Delivery"] is None and by_isbn[ISBNS[2]]["Used Delivery"] == "Tue, November 3 - Fri, November 6"
    assert by_isbn[ISBNS[4]]["New Delivery"] == "Mon, November 9 - Thu, November 12"
    last = ws.cell(row=1, column=ws.max_column)
    assert last.value == "Used Delivery" and "Order within" in last.comment.text

    up = upload(c, "\n".join(["ISBN", ISBNS[0], ISBNS[3], ISBNS[0]]))
    inv_id = up["job"]["id"]
    c.post(f"/api/inventory/jobs/{inv_id}/start", json={})
    assert wait_done(c, "inventory", inv_id)["status"] == "done"
    _, rows_out, _ = sheet(c, "inventory", inv_id)
    assert [(r["New Delivery"], r["Used Delivery"]) for r in rows_out] == [
        ("Thu, November 5 - Tue, November 10", "Tue, November 3 - Fri, November 6"),
        (None, None),  # sold out: no window
        ("Thu, November 5 - Tue, November 10", "Tue, November 3 - Fri, November 6"),  # duplicate reuses the answer
    ]


def test_v1_0_databases_are_upgraded_in_place(tmp_path):
    """A client PC upgrading from v1.0 keeps its runs; the new columns are added on start-up."""
    import sqlite3

    from awesomebooks_pro.storage import CatalogStore, InventoryStore

    for cls, name in ((CatalogStore, "catalog.db"), (InventoryStore, "inventory.db")):
        path = tmp_path / name
        store = cls(path)
        store.conn.execute("ALTER TABLE items DROP COLUMN new_delivery")
        store.conn.execute("ALTER TABLE items DROP COLUMN used_delivery")
        extra = {"url": "u", "section": "s"} if cls is CatalogStore else {"filename": "f.csv"}
        job_id = store.create_job(sessions=1, pace="balanced", **extra)
        store.insert_items(job_id, [{"isbn": ISBNS[0], "title": "Kept"}])
        store.close()

        reopened = cls(path)  # what the new version does when the app starts
        columns = {row[1] for row in sqlite3.connect(path).execute("PRAGMA table_info(items)")}
        assert {"new_delivery", "used_delivery"} <= columns
        item = reopened.claim_next(job_id)
        assert item["title"] == "Kept"
        reopened.complete_item(item["id"], {"new_delivery": "Thu, November 5 - Tue, November 10"})
        assert next(reopened.iter_items(job_id))["new_delivery"] == "Thu, November 5 - Tue, November 10"
        reopened.close()
