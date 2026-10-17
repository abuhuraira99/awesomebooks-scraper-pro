"""Excel exports.

One row per book (ISBN-13), with the New and the Used offer side by side:
available or not, the price a US shopper pays, and the exact stock. Only
columns the site actually provides are written. If a run skipped the detail
pass, the columns that pass would have filled are left out instead of being
exported empty, and a condition the listing didn't show stays blank (unknown),
never "No".
"""

from __future__ import annotations

import os
import re
import time
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

APP_NAME = "AwesomeBooks Scraper Pro"
QTY_NOTE = (
    "Exact stock from AwesomeBooks' own product data for this condition. Shoppers see '20+' on the page once stock is above 20."
)
PRICE_NOTE = (
    "What a US shopper pays on AwesomeBooks (shipping country United States, USD). The site states the price includes shipping."
)
AVAILABLE_NOTE = "Yes = this condition can be bought now. Blank = not checked (quick run without book pages)."

HEADER_FILL = PatternFill("solid", fgColor="123B2F")
NEW_FILL = PatternFill("solid", fgColor="1D5A46")
USED_FILL = PatternFill("solid", fgColor="7A5B1E")
HEADER_FONT = Font(bold=True, color="F7F4EC")
LINK_FONT = Font(color="0F6B4F", underline="single")
MONEY = '"$"#,##0.00'


@dataclass(frozen=True)
class Column:
    header: str
    width: int
    value: Callable[[dict[str, Any]], Any]
    fmt: str | None = None
    link: bool = False
    note: str | None = None
    fill: PatternFill | None = None


def _yes_no(v: Any) -> str | None:
    if v is None:
        return None
    return "Yes" if v else "No"


def _ts(v: Any) -> str | None:
    return datetime.fromtimestamp(v).strftime("%Y-%m-%d %H:%M:%S") if v else None


def _col(header: str, width: int, key: str, fmt: str | None = None, **kw: Any) -> Column:
    return Column(header, width, lambda r, k=key: r.get(k), fmt, **kw)


def _offer_columns(done: Callable[[dict[str, Any]], bool]) -> list[tuple[Column, bool]]:
    """New / Used blocks. ``done`` says whether a row has been checked (blank cells otherwise)."""

    def when(fn: Callable[[dict[str, Any]], Any]) -> Callable[[dict[str, Any]], Any]:
        return lambda r: fn(r) if done(r) else None

    out: list[tuple[Column, bool]] = []
    for prefix, label, fill in (("new", "New", NEW_FILL), ("used", "Used", USED_FILL)):
        out += [
            (
                Column(
                    f"{label} Available",
                    11,
                    when(lambda r, p=prefix: _yes_no(r.get(f"{p}_available"))),
                    note=AVAILABLE_NOTE,
                    fill=fill,
                ),
                False,
            ),
            (
                Column(
                    f"{label} Price (USD)", 13, when(lambda r, p=prefix: r.get(f"{p}_price")), MONEY, note=PRICE_NOTE, fill=fill
                ),
                False,
            ),
            (Column(f"{label} Qty", 9, when(lambda r, p=prefix: r.get(f"{p}_qty")), "0", note=QTY_NOTE, fill=fill), True),
        ]
    return out


def _catalog_done(r: dict[str, Any]) -> bool:
    return r.get("state") == "done"


def _catalog_status(r: dict[str, Any]) -> str | None:
    if r.get("state") == "error":
        return "Check failed"
    if r.get("state") != "done":
        return "Not checked"
    if r.get("new_available") is None and r.get("used_available") is None:
        return None
    return "In Stock" if r.get("new_available") or r.get("used_available") else "Out of Stock"


# (column, needs the detail pass): detail-only columns are left out of quick runs.
CATALOG_LAYOUT: list[tuple[Column, bool]] = [
    (_col("ISBN-13", 16, "isbn"), False),
    (_col("Title", 46, "title"), False),
    (_col("Author", 26, "author"), False),
    (_col("Format", 14, "format"), False),
    (Column("Availability", 13, _catalog_status), False),
    *_offer_columns(_catalog_done),
    (_col("RRP (USD)", 11, "rrp", MONEY), False),
    (_col("Publisher", 26, "publisher"), True),
    (_col("Publication Date", 15, "publication_date"), True),
    (_col("Pages", 8, "page_count", "0"), True),
    (_col("Dimensions", 20, "dimensions"), True),
    (_col("Weight", 10, "weight"), True),
    (_col("Category", 40, "category"), True),
    (_col("Note", 30, "note"), True),
    (_col("Product URL", 42, "product_url", link=True), False),
    (_col("Image URL", 42, "image_url", link=True), False),
]


def catalog_columns(with_details: bool) -> list[Column]:
    return [col for col, detail_only in CATALOG_LAYOUT if with_details or not detail_only]


def _inventory_done(r: dict[str, Any]) -> bool:
    return r.get("state") == "done"


def _inventory_status(r: dict[str, Any]) -> str:
    if r["state"] == "invalid":
        return "Invalid ISBN"
    if r["state"] == "error":
        return "Check failed"
    if r["state"] != "done":
        return "Not checked"
    return "In Stock" if r["stock_status"] == "in_stock" else "Out of Stock"


INVENTORY_COLUMNS: list[Column] = [
    Column("SKU", 22, lambda r: r["raw_sku"]),
    Column("ISBN-13", 16, lambda r: r["isbn"]),
    Column("Status", 13, _inventory_status),
    Column(
        "Quantity",
        10,
        lambda r: r["quantity"] if r["state"] == "done" else None,
        "0",
        note="Units in stock across New and Used. 0 = out of stock or not listed. " + QTY_NOTE,
    ),
    *[col for col, _ in _offer_columns(_inventory_done)],
    Column("RRP (USD)", 11, lambda r: r.get("rrp"), MONEY),
    Column("Title", 44, lambda r: r.get("title")),
    Column("Author", 24, lambda r: r.get("author")),
    Column("Format", 13, lambda r: r.get("format")),
    Column("Listed on AwesomeBooks", 21, lambda r: _yes_no(r["found"]) if r["state"] == "done" else None),
    Column("Note", 32, lambda r: r.get("note") or (r.get("error") if r["state"] in ("invalid", "error") else None)),
    Column("Product URL", 42, lambda r: r.get("product_url"), link=True),
    Column("Checked At", 19, lambda r: _ts(r.get("checked_at"))),
]


MAX_HYPERLINKS = 60_000  # Excel refuses more than ~66k hyperlinks per sheet; beyond this URLs stay plain text


def write_workbook(path: Path, sheet_title: str, columns: list[Column], rows: Iterable[dict[str, Any]]) -> Path:
    """Stream rows into a styled workbook. Memory stays flat, so 100k+ rows are fine."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet_title[:31])
    ws.freeze_panes = "C2"  # write-only sheets need layout settings before the first row
    for idx, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = col.width
    ws.row_dimensions[1].height = 30

    header = []
    for col in columns:
        cell = WriteOnlyCell(ws, value=col.header)
        cell.fill = col.fill or HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if col.note:
            cell.comment = Comment(col.note, APP_NAME)
        header.append(cell)
    ws.append(header)

    links = 0
    count = 1
    for row in rows:
        out = []
        for col in columns:
            value = col.value(row)
            cell = WriteOnlyCell(ws, value=value)
            if (
                isinstance(value, str)
                and value.startswith(("=", "+", "-", "@"))
                and not value.lstrip("+-").replace(".", "", 1).isdigit()
            ):
                cell.data_type = "s"  # never let scraped or uploaded text become an Excel formula
            if col.fmt and value is not None:
                cell.number_format = col.fmt
            if col.link and value and links < MAX_HYPERLINKS:
                cell.hyperlink = value
                cell.font = LINK_FONT
                links += 1
            out.append(cell)
        ws.append(out)
        count += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{count}"
    wb.properties.creator = APP_NAME
    wb.properties.created = datetime.fromtimestamp(time.time())
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Destinations
# ---------------------------------------------------------------------------


def downloads_dir() -> Path:
    """The user's real Downloads folder (honours OneDrive / relocated folders on Windows)."""
    if os.name == "nt":
        try:
            import ctypes
            import uuid as _uuid
            from ctypes import wintypes

            class GUID(ctypes.Structure):
                _fields_ = [("d1", wintypes.DWORD), ("d2", wintypes.WORD), ("d3", wintypes.WORD), ("d4", ctypes.c_ubyte * 8)]

            raw = _uuid.UUID("{374DE290-123F-4565-9164-39C4925E467B}").bytes_le  # FOLDERID_Downloads
            guid = GUID.from_buffer_copy(raw)
            out = ctypes.c_wchar_p()
            if ctypes.windll.shell32.SHGetKnownFolderPath(ctypes.byref(guid), 0, None, ctypes.byref(out)) == 0:
                path = Path(out.value)
                ctypes.windll.ole32.CoTaskMemFree(out)
                if path.is_dir():
                    return path
        except Exception:
            pass
    home = Path.home() / "Downloads"
    return home if home.is_dir() else Path.home()


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60] or "section"


def export_filename(kind: str, job: dict[str, Any]) -> str:
    stamp = time.strftime("%Y%m%d-%H%M")
    if kind == "catalog":
        return f"awesomebooks-{_slug(job['section'])}-{stamp}.xlsx"
    base = re.sub(r"[^\w\- ]+", "_", Path(job["filename"]).stem).strip() or "stock"
    return f"{base}-stock-check-{stamp}.xlsx"


def export_job(kind: str, store: Any, job: dict[str, Any], path: Path) -> Path:
    if kind == "catalog":
        return write_workbook(path, "Catalog", catalog_columns(bool(job["with_details"])), store.iter_items(job["id"]))
    return write_workbook(path, "Stock Check", INVENTORY_COLUMNS, store.iter_items(job["id"]))


def unique_path(folder: Path, name: str) -> Path:
    path = folder / name
    n = 2
    while path.exists():
        path = folder / f"{Path(name).stem} ({n}){Path(name).suffix}"
        n += 1
    return path
